import openpyxl
from pathlib import Path

file_path = Path(r"C:\Users\User\OneDrive\Documents\Commission\7. Presentation\Finance Output\Commission_Pack_2026_20260625_083605.xlsx")
if not file_path.exists():
    # Find latest file
    folder = Path(r"C:\Users\User\OneDrive\Documents\Commission\7. Presentation\Finance Output")
    files = list(folder.glob("*.xlsx"))
    if files:
        file_path = sorted(files)[-1]
        print(f"Using latest file: {file_path}")
    else:
        print("No Excel files found.")
        exit(1)

wb = openpyxl.load_workbook(file_path, read_only=True)
print(f"Loaded Excel: {file_path.name}")
print("Sheets:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    found = False
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        row_str = " ".join(str(cell) for cell in row if cell is not None)
        if "hosaini" in row_str.lower():
            if not found:
                print(f"\nMatches in sheet: {sheet_name}")
                found = True
            print(f"  Line {r_idx}: {row_str}")
