import openpyxl
from pathlib import Path
import sys

sys.stdout.reconfigure(encoding='utf-8')

def main():
    folder = Path(r"C:\Users\User\OneDrive\Documents\Commission\7. Presentation\Finance Output")
    files = sorted(folder.glob("Outsource_Commission_Pack_2026_*.xlsx"))
    if not files:
        print("No output Excel files found!")
        return
        
    latest_file = files[-1]
    print(f"Dumping Basic - By Invoice from file: {latest_file.name}")
    print("=" * 80)
    
    wb = openpyxl.load_workbook(latest_file, data_only=True)
    if "Basic - By Invoice" in wb.sheetnames:
        sheet = wb["Basic - By Invoice"]
        print(f"Total rows in sheet: {sheet.max_row}")
        for r in range(1, sheet.max_row + 1):
            row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
            if any(row_vals):
                print(f"Row {r}: {row_vals}")
    else:
        print("Basic - By Invoice sheet not found!")

if __name__ == '__main__':
    main()
