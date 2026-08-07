import openpyxl
from pathlib import Path
import sys

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    file_path = Path(r"1. Basic Commission/1. Excel/3. Book (Basic & NFP Refrence from Finance).xlsx")
    if not file_path.exists():
        print("File not found!")
        return
    wb = openpyxl.load_workbook(file_path, data_only=True)
    if "Paid on Jun" not in wb.sheetnames:
        print("Paid on Jun sheet not found!")
        return
        
    ws = wb["Paid on Jun"]
    print(f"Paid on Jun sheet (Total rows: {ws.max_row}):")
    print("=" * 100)
    for r_idx in range(1, ws.max_row + 1):
        row_vals = [cell.value for cell in ws[r_idx]]
        if any(row_vals):
            print(f"Row {r_idx}: {[str(v) if v is not None else '' for v in row_vals]}")

if __name__ == '__main__':
    main()
