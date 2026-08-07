import openpyxl
from pathlib import Path

def main():
    excel_path = Path("C:/Users/User/OneDrive/Documents/Commission/2. NFP Commission/1. Excel/1. Net Floor Price Commission.xlsx")
    if not excel_path.exists():
        print(f"File not found: {excel_path}")
        return
        
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    for name in wb.sheetnames:
        if name == "Condition":
            continue
        ws = wb[name]
        print(f"\n================ SHEET: {name} ================")
        
        rows = list(ws.iter_rows(values_only=True))
        for idx, r in enumerate(rows[:15]):
            clean_r = []
            for x in r:
                if x is None:
                    clean_r.append("")
                else:
                    # Keep only ascii
                    s = str(x).encode('ascii', errors='ignore').decode('ascii')
                    clean_r.append(s)
            
            # Print only if there is something non-empty in the first 15 columns
            if any(c != "" for c in clean_r[:15]):
                print(f"Row {idx+1:02d}: {clean_r[:15]}")

if __name__ == '__main__':
    main()
