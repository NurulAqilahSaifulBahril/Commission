import openpyxl
from pathlib import Path
import sys

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    file_path = Path(r"C:\Users\User\OneDrive\Documents\Commission\7. Presentation\Finance Output\Commission_Pack_2026_20260702_162015.xlsx")
    if not file_path.exists():
        print(f"File not found: {file_path}")
        return
        
    wb = openpyxl.load_workbook(file_path, data_only=True)
    print("Sheets in workbook:")
    for name in wb.sheetnames:
        print(f"  - {name}")
        
    print("\nSearching for 'ching chin min', '1008343' in all sheets:")
    found = False
    for name in wb.sheetnames:
        ws = wb[name]
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            row_str = " ".join(str(cell) for cell in row if cell is not None)
            if any(term in row_str.lower() for term in ['ching chin min', '1008343']):
                print(f"[{name}] Row {r_idx}: {[str(c) for c in row if c is not None]}")
                found = True
    if not found:
        print("No matches found in the entire workbook.")

if __name__ == "__main__":
    main()
