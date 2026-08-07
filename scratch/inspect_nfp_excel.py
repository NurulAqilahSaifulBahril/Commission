import openpyxl
from pathlib import Path

def main():
    excel_path = Path("C:/Users/User/OneDrive/Documents/Commission/2. NFP Commission/1. Excel/1. Net Floor Price Commission.xlsx")
    if not excel_path.exists():
        print(f"File not found: {excel_path}")
        return
        
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    print("Sheets in workbook:")
    for sheet in wb.sheetnames:
        print(f" - {sheet}")
        
    for name in wb.sheetnames:
        print(f"\n--- First 10 rows of sheet: {name} ---")
        ws = wb[name]
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx >= 10:
                break
            print(row)

if __name__ == '__main__':
    main()
