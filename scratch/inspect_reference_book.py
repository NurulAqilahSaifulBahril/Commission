import openpyxl
from pathlib import Path
import sys

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    file_path = Path(r"1. Basic Commission/1. Excel/3. Book (Basic & NFP Refrence from Finance).xlsx")
    if not file_path.exists():
        print("File not found!")
        return
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    target_names = [
        "hosaini", "hong chin jia", "cho chooi peng", "prema", "mandave", 
        "loh kah guan", "aw mun lai", "ho lee"
    ]
    
    print("Searching for target names in the reference book:")
    for name in wb.sheetnames:
        ws = wb[name]
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            row_str = " ".join(str(cell).lower() for cell in row if cell is not None)
            for target in target_names:
                if target in row_str:
                    print(f"[{name}] Row {r_idx} ({target}): {[str(c) for c in row if c is not None]}")

if __name__ == '__main__':
    main()
