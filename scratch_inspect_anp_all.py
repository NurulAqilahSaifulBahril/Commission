import openpyxl
from pathlib import Path
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_path = Path(r"C:\Users\User\OneDrive\Documents\Commission\7. Presentation\Finance Output\Commission_Pack_2026_20260625_083605.xlsx")
wb = openpyxl.load_workbook(file_path, data_only=True)

for sheet_name in ['ANP - By Agent', 'ANP - By Invoice']:
    ws = wb[sheet_name]
    print(f"\nSheet: {sheet_name}")
    for r_idx in range(1, 6):
        row_vals = [cell.value for cell in ws[r_idx]]
        print(f"  Row {r_idx}: {row_vals}")
    ws = wb['ANP - By Invoice']
    print(f"\nLines 60 to 75 of ANP - By Invoice:")
    for r in range(60, 76):
        row_vals = [cell.value for cell in ws[r]]
        print(f"  Row {r}: {row_vals}")
