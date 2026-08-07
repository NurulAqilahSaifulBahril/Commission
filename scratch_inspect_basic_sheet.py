import openpyxl
from pathlib import Path
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_path = Path(r"C:\Users\User\OneDrive\Documents\Commission\7. Presentation\Finance Output\Commission_Pack_2026_20260625_083605.xlsx")
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb['Basic - Residential & Shop Lot']

# Print header rows (first 5 rows)
for r in range(1, 6):
    row_vals = [cell.value for cell in ws[r]]
    print(f"Header Row {r}: {row_vals}")

# Print lines 134 and 135
print("\nLine 134:")
print([cell.value for cell in ws[134]])
print("\nLine 135:")
print([cell.value for cell in ws[135]])
